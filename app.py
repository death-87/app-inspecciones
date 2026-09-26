import io
import base64
import requests
import streamlit as st
import pandas as pd
import gspread
import plotly.express as px
from google.oauth2.service_account import Credentials
from datetime import datetime

# Librerías para generación de PDF
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader

# Librerías para generación de Word
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# =========================================================
# CONFIGURACIÓN DE LA PÁGINA
# =========================================================
st.set_page_config(
    page_title="Sistema de Gestión de Activos Físicos",
    page_icon="☁️",
    layout="wide"
)

# Estilo compacto. El tema de .streamlit/config.toml define los colores base.
st.markdown("""
<style>
.block-container { max-width: 1240px; padding-top: 2.5rem; padding-bottom: 1rem; }
[data-testid="stSidebar"] { background: #edf2f6; border-right: 1px solid #dbe3eb; }
[data-testid="stSidebar"] [data-testid="stVerticalBlock"] { gap: .65rem; }
[data-testid="stVerticalBlock"] { gap: .75rem; }
[data-testid="stForm"] { background: #ffffff; border: 1px solid #dbe3eb; border-radius: 12px; padding: 1rem; }
h1, h2, h3 { color: #17324d; letter-spacing: -.025em; }
h3 { font-size: 1.2rem !important; padding-top: .2rem !important; }
.app-heading { border-left: 4px solid #087f73; padding: .1rem 0 .1rem 1rem; margin-bottom: .6rem; }
.app-heading h1 { font-size: 1.8rem; line-height: 1.2; margin: 0; padding: 0; }
.app-heading p { color: #526579; font-size: .9rem; margin: .35rem 0 0; }
.app-eyebrow { color: #087f73; font-size: .72rem; font-weight: 700; letter-spacing: .14em; margin-bottom: .35rem; }
.app-footer { border-top: 1px solid #dbe3eb; padding-top: .65rem; margin-top: 1rem; color: #526579; font-size: .75rem; }
[data-testid="stButton"] button, [data-testid="stFormSubmitButton"] button { border-radius: 8px; }
@media (max-width: 640px) {
    .block-container { padding: 2rem 1rem 1rem; }
    .app-heading h1 { font-size: 1.45rem; }
    [data-testid="stHorizontalBlock"] { flex-wrap: wrap; }
    [data-testid="stColumn"] { min-width: 100% !important; flex: 1 1 100% !important; }
}
</style>
""", unsafe_allow_html=True)

# 🔗 URLs RAW DE LOGO, FRANJA Y PERSONAJE EN GITHUB
URL_LOGO_GITHUB = "https://raw.githubusercontent.com/death-87/app-inspecciones/main/logo.png"
URL_FRANJA_GITHUB = "https://raw.githubusercontent.com/death-87/app-inspecciones/main/franja.png"
URL_PERSONAJE_GITHUB = "https://raw.githubusercontent.com/death-87/app-inspecciones/main/personaje.png"

# 🆔 ID DE TU HOJA DE GOOGLE SHEETS
SPREADSHEET_ID = "1eJpQXWqe4AyyrFm_6wlnfzm-KYSGPeTtX_EWCIJYE1I"

# =========================================================
# GESTIÓN DE ROLES INTERNOS Y SESIÓN
# =========================================================
USUARIOS_SISTEMA = {'jnavarrete': {'password': 'Mechanix123', 'rol': 'operador', 'nombre': 'jnavarrete (Agregar Datos)'}}

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False
if "usuario_actual" not in st.session_state:
    st.session_state.usuario_actual = None
if "rol_actual" not in st.session_state:
    st.session_state.rol_actual = "invitado"
if "nombre_usuario" not in st.session_state:
    st.session_state.nombre_usuario = "Visitante"

# =========================================================
# CACHÉ DE IMÁGENES
# =========================================================
@st.cache_data(ttl=3600)
def obtener_bytes_imagen(url):
    try:
        response = requests.get(url, timeout=5)
        if response.status_code == 200:
            return response.content
    except Exception:
        pass
    return None

def obtener_bytes_logo():
    return obtener_bytes_imagen(URL_LOGO_GITHUB)

def obtener_bytes_franja():
    return obtener_bytes_imagen(URL_FRANJA_GITHUB)

def obtener_bytes_personaje():
    return obtener_bytes_imagen(URL_PERSONAJE_GITHUB)

# =========================================================
# CONEXIÓN CON GOOGLE SHEETS (MEDIANTE SERVICE ACCOUNT)
# =========================================================
@st.cache_resource
def conectar_google_sheets():
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ]
    creds_dict = dict(st.secrets["connections"]["gsheets"])
    if "private_key" in creds_dict:
        creds_dict["private_key"] = creds_dict["private_key"].replace("\\n", "\n")

    credentials = Credentials.from_service_account_info(creds_dict, scopes=scopes)
    client = gspread.authorize(credentials)
    return client.open_by_key(SPREADSHEET_ID)

def obtener_hoja_actividades():
    client = conectar_google_sheets()
    return client.sheet1

def obtener_hoja_planificacion():
    client = conectar_google_sheets()
    try:
        return client.worksheet("Planificacion")
    except gspread.exceptions.WorksheetNotFound:
        ws = client.add_worksheet(title="Planificacion", rows=100, cols=25)
        ws.append_row([
            "planta", "tipo_informe", "circuito_equipo", "fecha_inicio", "fecha_fin", 
            "contador_inspeccion", "programa", "inspector_dci", "inspector_ingemars", 
            "dias_ing_enap", "otep", "informe_iv", "status", "fecha_entrega_ope", 
            "contador_liberacion", "liberacion_final", "contador_enap", 
            "contador_cumplimiento_enap", "contador_dias_informe", "archivo_url", "observaciones"
        ])
        return ws

def cargar_datos_sheets():
    try:
        sheet = obtener_hoja_actividades()
        filas = sheet.get_all_values()
        if len(filas) <= 1:
            return pd.DataFrame(columns=[
                "fecha", "semana", "planta", "inspector", "tag_equipo", 
                "actividad_realizada", "avance", "observaciones", "estado_liberacion"
            ])
        encabezados = [c.strip().lower() for c in filas[0]]
        return pd.DataFrame(filas[1:], columns=encabezados)
    except Exception as e:
        st.error(f"Error al leer Google Sheets: {e}")
        return pd.DataFrame()

def cargar_datos_planificacion():
    try:
        sheet = obtener_hoja_planificacion()
        filas = sheet.get_all_values()
        if len(filas) <= 1:
            return pd.DataFrame()
        encabezados = [c.strip().lower() for c in filas[0]]
        return pd.DataFrame(filas[1:], columns=encabezados)
    except Exception:
        return pd.DataFrame()

def obtener_ultimo_registro_tag(tag_busqueda):
    df_historial = cargar_datos_sheets()
    if df_historial.empty or "tag_equipo" not in df_historial.columns:
        return None
    df_tag = df_historial[df_historial['tag_equipo'].astype(str).str.upper() == tag_busqueda.strip().upper()]
    if not df_tag.empty:
        return df_tag.iloc[-1].to_dict()
    return None

# =========================================================
# DIBUJO Y EXPORTACIÓN DE REPORTES (PDF Y WORD)
# =========================================================
def dibujar_plantilla(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(colors.HexColor("#f8faf6"))
    canvas.rect(0, 0, letter[0], letter[1], fill=1, stroke=0)
    canvas.setFillColor(colors.HexColor("#619b40"))
    canvas.rect(0, letter[1] - (2 * cm), letter[0], 2 * cm, fill=1, stroke=0)

    logo_bytes = obtener_bytes_logo()
    if logo_bytes:
        try:
            img_stream = io.BytesIO(logo_bytes)
            img = ImageReader(img_stream)
            canvas.drawImage(img, 30, letter[1] - (2 * cm) - 55, width=120, height=45, preserveAspectRatio=True, mask='auto')
        except Exception:
            pass

    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#6B7280"))
    canvas.drawCentredString(letter[0] / 2.0, 26, "SERVICIO DE INSPECCIÓN Y EVALUACIÓN DE ACTIVOS FÍSICOS DE ENAP REFINERÍAS S.A.")
    canvas.drawCentredString(letter[0] / 2.0, 16, "CONTRATO N° AC 31104857")
    canvas.drawRightString(letter[0] - 30, 16, f"Pág. {canvas.getPageNumber()}")
    canvas.restoreState()

def generar_pdf_informe(df_filtrado, titulo_doc, subtitulo_doc):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=4.5 * cm, bottomMargin=2.2 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#1E3A8A'), alignment=1, spaceAfter=4)
    subtitle_style = ParagraphStyle('DocSubTitle', parent=styles['Normal'], fontSize=11, leading=14, textColor=colors.HexColor('#4B5563'), alignment=1, spaceAfter=14)
    inspector_heading_style = ParagraphStyle('InspectorHeader', parent=styles['Heading2'], fontSize=11, leading=14, textColor=colors.HexColor('#1E3A8A'), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6)
    cell_header_style = ParagraphStyle('HeaderStyle', parent=styles['Normal'], fontSize=8.5, leading=10, textColor=colors.white, fontName='Helvetica-Bold', alignment=1)
    cell_body_style = ParagraphStyle('BodyStyle', parent=styles['Normal'], fontSize=8, leading=10, textColor=colors.HexColor('#1F2937'))

    story = [Paragraph(titulo_doc, title_style), Paragraph(f"<b>{subtitulo_doc}</b>", subtitle_style), Spacer(1, 6)]

    if not df_filtrado.empty:
        inspectores_grupos = df_filtrado.groupby('inspector', sort=False)
        for idx, (inspector_nom, group_df) in enumerate(inspectores_grupos):
            story.append(Paragraph(f"👷‍♂️ Inspector: <b>{inspector_nom}</b>", inspector_heading_style))
            table_data = [[
                Paragraph("Fecha / Planta", cell_header_style),
                Paragraph("Actividad Realizada", cell_header_style),
                Paragraph("Avance", cell_header_style),
                Paragraph("Estado", cell_header_style),
                Paragraph("Observaciones", cell_header_style)
            ]]
            for _, row in group_df.iterrows():
                table_data.append([
                    Paragraph(f"<b>Fecha:</b> {row.get('fecha', '-')}<br/><b>Planta:</b> {row.get('planta', '-')}<br/><b>TAG:</b> {row.get('tag_equipo', '-')}", cell_body_style),
                    Paragraph(str(row.get('actividad_realizada', '-')), cell_body_style),
                    Paragraph(str(row.get('avance', '-')), cell_body_style),
                    Paragraph(str(row.get('estado_liberacion', '-')), cell_body_style),
                    Paragraph(str(row.get('observaciones', '-')) if row.get('observaciones') else "-", cell_body_style)
                ])
            t = Table(table_data, colWidths=[110, 150, 50, 92, 150])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E3A8A')),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#D1D5DB')),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F3F4F6')]),
                ('PADDING', (0,0), (-1,-1), 5),
            ]))
            story.append(t)
            story.append(Spacer(1, 10))

    doc.build(story, onFirstPage=dibujar_plantilla, onLaterPages=dibujar_plantilla)
    buffer.seek(0)
    return buffer

def generar_word_informe(df_filtrado, titulo_doc, subtitulo_doc):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(1.0)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(titulo_doc)
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)

    if not df_filtrado.empty:
        for inspector_nom, group_df in df_filtrado.groupby('inspector', sort=False):
            insp_p = doc.add_paragraph()
            insp_run = insp_p.add_run(f"Inspector: {inspector_nom}")
            insp_run.font.bold = True
            insp_run.font.size = Pt(11)

            table = doc.add_table(rows=1, cols=5)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            headers = ["Fecha / Planta", "Actividad Realizada", "Avance", "Estado", "Observaciones"]
            for i, header_text in enumerate(headers):
                hdr_cells[i].text = header_text
                hdr_cells[i].paragraphs[0].runs[0].font.bold = True

            for _, row in group_df.iterrows():
                row_cells = table.add_row().cells
                row_cells[0].text = f"Fecha: {row.get('fecha', '-')}\nPlanta: {row.get('planta', '-')}\nTAG: {row.get('tag_equipo', '-')}"
                row_cells[1].text = str(row.get("actividad_realizada", "-"))
                row_cells[2].text = str(row.get("avance", "-"))
                row_cells[3].text = str(row.get("estado_liberacion", "-"))
                row_cells[4].text = str(row.get("observaciones", "-")) if row.get("observaciones") else "-"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# =========================================================
# LISTAS Y CONSTANTES
# =========================================================
LISTA_INSPECTORES = ["Juan Navarrete", "Jorge Hernandez", "Arlem Sarmiento", "Harold Castillo", "Miguel Chirinos"]
LISTA_INSPECTORES_DCI = ["Eduardo Delgado", "Felipe Ponce", "José de la Cruz", "Pablo Ruiz", "Luis Durán"]
ESTADOS_STATUS = ["Finalizado", "En curso", "Pendiente"]
TIPOS_INFORME = ["IV", "Ensayo Dureza", "Ensayo LP", "IHV", "Boroscopio"]
LISTA_PLANTAS = ["A0AEX", "A0ALQ", "A0BUT", "A0CCK", "A0CCR", "A0CKR", "A0HDG", "A0HDT", "A0HCK", "A0ISO", "A0LAB", "A0MHC", "A0NHT", "A0SAR", "A0SHP", "A0SWS", "AACID", "AAMAR", "AAMIN", "AAMPL", "AANTO", "AAREF", "AASER", "AALQU", "ADESO", "ADEV1", "ADEV2", "ADIPE", "AE501", "ALNHT", "ALPG1", "ALPG2", "ALPG3", "AMACO", "AMDEA", "AMRX1", "AMRX2", "AMRX3", "AMRX4", "AMVPR", "AOLEO", "APBMP", "APBTQ", "APCAR", "APFEN", "APRCO", "ARPLU", "AREFO", "AREMO", "ARILE", "ASAIC", "ASOLV", "ASPLI", "ASRCO", "ASUEL", "ASVAQ", "ASVAP", "ASWS2", "ASYBR", "ASEFL", "ASEFQ", "ATOP1", "ATOP2", "ATRAG", "AURA1", "AURA2", "AURA3", "AVAC1", "AVAC2", "ACOKE"]
LISTA_SEMANAS = [f"Semana {i}" for i in range(1, 53)]
ESTADOS_LIBERACION = ["Liberado / Conforme (Aprobado)", "Pendiente de Reparación", "Rechazado", "En Proceso de Inspección", "En Espera de END / Pruebas"]

semana_actual_num = datetime.now().isocalendar()[1]
idx_semana_defecto = max(0, min(semana_actual_num - 1, len(LISTA_SEMANAS) - 1))

# =========================================================
# BARRA LATERAL (LOGIN INTERNO Y CONTROL DE ACCESO)
# =========================================================
logo_bytes_sidebar = obtener_bytes_logo()
if logo_bytes_sidebar:
    st.sidebar.image(logo_bytes_sidebar, width=145)

# Invalidar también sesiones anteriores de cuentas que ya no tienen acceso.
if st.session_state.usuario_actual not in USUARIOS_SISTEMA:
    st.session_state.autenticado = False
    st.session_state.usuario_actual = None
    st.session_state.rol_actual = "invitado"
    st.session_state.nombre_usuario = "Visitante"

st.sidebar.markdown("### Acceso privado")
if not st.session_state.autenticado:
    with st.sidebar.form("form_login"):
        user_input = st.text_input("Usuario:")
        pass_input = st.text_input("Contraseña:", type="password")
        btn_login = st.form_submit_button("Iniciar sesión", type="primary", use_container_width=True)
        if btn_login:
            if user_input in USUARIOS_SISTEMA and USUARIOS_SISTEMA[user_input]["password"] == pass_input:
                st.session_state.autenticado = True
                st.session_state.usuario_actual = user_input
                st.session_state.rol_actual = USUARIOS_SISTEMA[user_input]["rol"]
                st.session_state.nombre_usuario = USUARIOS_SISTEMA[user_input]["nombre"]
                st.rerun()
            else:
                st.sidebar.error("❌ Usuario o contraseña incorrectos")
    st.sidebar.info("ℹ️ Inicia sesión con tu cuenta para acceder a la aplicación.")
else:
    st.sidebar.caption(f"Sesión activa · {st.session_state.usuario_actual}")
    if st.sidebar.button("Cerrar sesión", use_container_width=True):
        st.session_state.autenticado = False
        st.session_state.usuario_actual = None
        st.session_state.rol_actual = "invitado"
        st.session_state.nombre_usuario = "Visitante"
        st.rerun()

# No renderizar módulos ni consultar Google Sheets sin una sesión autorizada.
if not st.session_state.autenticado or st.session_state.usuario_actual not in USUARIOS_SISTEMA:
    st.info("🔐 Acceso privado. Inicia sesión en la barra lateral.")
    st.stop()

st.sidebar.markdown("---")

# =========================================================
# ENCABEZADO PRINCIPAL
# =========================================================
st.markdown("""
<div class="app-heading">
    <div class="app-eyebrow">CONTROL OPERATIVO · QA/QC</div>
    <h1>Gestión de activos físicos</h1>
    <p>Inspecciones, historial y planificación de operaciones.</p>
</div>
""", unsafe_allow_html=True)

# =========================================================
# MENÚ NAVEGACIÓN PRINCIPAL
# =========================================================
menu = st.sidebar.radio(
    "Navegación",
    [
        "Registrar actividad", 
        "Historial e informes", 
        "Planificación"
    ]
)

st.sidebar.caption("Gestión de inspecciones · ENAP")

rol_usuario = st.session_state.rol_actual

# =========================================================
# MÓDULO 1: REGISTRO DE ACTIVIDADES
# =========================================================
if menu == "Registrar actividad":
    st.subheader("Registrar actividad")
    if rol_usuario == "invitado":
        st.warning("⚠️ Tu cuenta actual es de **Visitante (Solo Lectura)**. Inicia sesión en la barra lateral para registrar actividades.")

    with st.expander("Retomar una actividad por TAG", expanded=False):
        tag_para_retomar = st.text_input("TAG del equipo", placeholder="Ej: C-1302")
        btn_cargar_tag = st.button("Cargar último registro")

    def_inspector = LISTA_INSPECTORES[0]
    def_planta = LISTA_PLANTAS[0]
    def_tag = ""
    def_avance = 0
    def_actividad = ""
    def_obs = ""
    def_estado = ESTADOS_LIBERACION[3]

    if btn_cargar_tag and tag_para_retomar:
        ultimo_reg = obtener_ultimo_registro_tag(tag_para_retomar)
        if ultimo_reg:
            avance_str = str(ultimo_reg.get("avance", "0")).replace("%", "").strip()
            avance_val = int(float(avance_str)) if avance_str.replace('.', '', 1).isdigit() else 0
            st.success(f"✅ Último registro cargado para **{tag_para_retomar}** (Avance previo: {avance_val}%).")
            def_tag = str(ultimo_reg.get("tag_equipo", tag_para_retomar))
            def_actividad = str(ultimo_reg.get("actividad_realizada", ""))
            def_obs = f"Continuación de inspección anterior. Obs previas: {ultimo_reg.get('observaciones', '')}"
            def_avance = avance_val
            if ultimo_reg.get("inspector") in LISTA_INSPECTORES:
                def_inspector = ultimo_reg.get("inspector")
            if ultimo_reg.get("planta") in LISTA_PLANTAS:
                def_planta = ultimo_reg.get("planta")

    with st.form("form_actividades_inspector", clear_on_submit=True):
        idx_insp = LISTA_INSPECTORES.index(def_inspector) if def_inspector in LISTA_INSPECTORES else 0
        idx_plan = LISTA_PLANTAS.index(def_planta) if def_planta in LISTA_PLANTAS else 0
        idx_est = ESTADOS_LIBERACION.index(def_estado) if def_estado in ESTADOS_LIBERACION else 0
        col1, col2, col3 = st.columns([2, 1, 1])
        with col1:
            inspector_seleccionado = st.selectbox("Inspector asignado", LISTA_INSPECTORES, index=idx_insp)
        with col2:
            fecha_actividad = st.date_input("Fecha de inspección", datetime.now())
        with col3:
            semana_seleccionada = st.selectbox("Semana operativa", LISTA_SEMANAS, index=idx_semana_defecto)
        col4, col5, col6 = st.columns([1, 1, 2])
        with col4:
            planta_seleccionada = st.selectbox("Planta / unidad", LISTA_PLANTAS, index=idx_plan)
        with col5:
            tag_equipo = st.text_input("TAG del equipo", value=def_tag, placeholder="Ej: C-1302")
        with col6:
            estado_liberacion = st.selectbox("Estado de inspección", ESTADOS_LIBERACION, index=idx_est)
        porcentaje_avance = st.slider("Avance acumulado", min_value=0, max_value=100, value=def_avance, step=5, format="%d%%")
        col7, col8 = st.columns(2)
        with col7:
            actividad_realizada = st.text_area("Actividad realizada", value=def_actividad, height=100)
        with col8:
            observaciones = st.text_area("Observaciones", value=def_obs, height=100)
        btn_guardar = st.form_submit_button("Guardar actividad", type="primary")

        if btn_guardar:
            if rol_usuario == "invitado":
                st.error("❌ Acción no permitida para el rol de Visitante.")
            elif tag_equipo and actividad_realizada:
                try:
                    sheet = obtener_hoja_actividades()
                    nueva_fila = [
                        str(fecha_actividad),
                        semana_seleccionada,
                        planta_seleccionada,
                        inspector_seleccionado,
                        tag_equipo.strip(),
                        actividad_realizada.strip(),
                        f"{porcentaje_avance}%",
                        observaciones.strip(),
                        estado_liberacion
                    ]
                    sheet.append_row(nueva_fila)
                    st.cache_data.clear()
                    st.success(f"✅ ¡Nuevo registro de **{inspector_seleccionado}** guardado con éxito!")
                except Exception as ex:
                    st.error(f"❌ Ocurrió un error al guardar: {ex}")

# =========================================================
# MÓDULO 2: HISTORIAL E INFORMES
# =========================================================
elif menu == "Historial e informes":
    st.subheader("Historial e informes")
    df_historial = cargar_datos_sheets()
    if not df_historial.empty:
        st.caption(f"{len(df_historial):,} registros disponibles")
        st.dataframe(df_historial, use_container_width=True, hide_index=True)
        csv_data = df_historial.to_csv(index=False).encode('utf-8')
        st.download_button("📥 Descargar Tabla a CSV", csv_data, "historial.csv", "text/csv")
    else:
        st.info("ℹ️ No hay registros guardados.")

# =========================================================
# MÓDULO 3: REPORTE PLANIFICACIÓN
# =========================================================
elif menu == "Planificación":
    st.subheader("Planificación de operaciones")
    df_plan = cargar_datos_planificacion()
    if not df_plan.empty:
        st.caption(f"{len(df_plan):,} registros disponibles")
        st.dataframe(df_plan, use_container_width=True, hide_index=True)
    else:
        st.info("ℹ️ Sin datos de planificación.")

# PIE DE PÁGINA
st.markdown('<div class="app-footer">QA/QC · Servicio de inspección y evaluación de activos físicos</div>', unsafe_allow_html=True)
